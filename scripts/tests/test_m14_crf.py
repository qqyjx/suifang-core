#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""M14 智能 CRF 表单 (方案 §2.1)。

三组最要紧的断言:
  1. 逻辑引擎 —— 隐藏的题不参与必填校验(否则表单永远交不了, 且报错指向看不见的题),
     隐藏的题的旧答案不参与后续条件求值(否则产生幽灵数据)
  2. 版本管理 —— "修改不影响已有数据": 破坏性改动 + 已有填报 => 自动开新版
  3. Excel 推断的每一列都要说清"凭什么这么判"
"""
import os, sys, json, io
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _harness import hs, check, section, sub, finish, ensure_all_tables, db

PFX = 'T14'
ensure_all_tables()
items = lambda d: [it for _, it in hs._crf_items(d)]
def clean():
    db(("DELETE FROM platform_crf_response WHERE crf_code LIKE %s", (PFX+'%',)),
       ("DELETE FROM platform_crf WHERE code LIKE %s", (PFX+'%',)),
       ("DELETE FROM platform_patient WHERE patient_no LIKE %s", (PFX+'%',)))
clean()
db(("INSERT INTO platform_patient (patient_no,name) VALUES (%s,'CRF测试患者')", (PFX+'001',)))

section('1. 全题型覆盖 (方案 §2.1(1) 逐个点名的题型)')
YN=[{'label':'是','value':1},{'label':'否','value':0}]
DEF={'sections':[
    {'name':'基本','items':[
        {'id':'note1','type':'note','text':'以下由医护填写'},
        {'id':'nm','type':'text','text':'姓名','required':True},
        {'id':'idc','type':'text','text':'身份证号','format':'id_card'},
        {'id':'age','type':'number','text':'年龄','min':0,'max':130},
        {'id':'dt','type':'date','text':'访视日期','required':True},
        {'id':'desc','type':'paragraph','text':'主诉'},
        {'id':'sex','type':'single','text':'性别','options':[{'label':'男','value':1},{'label':'女','value':2}]},
        {'id':'sym','type':'multi','text':'症状','options':[{'label':'头痛','value':1},{'label':'失眠','value':2},{'label':'乏力','value':3}]},
        {'id':'dept','type':'select','text':'科室','options':[{'label':'内科','value':1},{'label':'外科','value':2}]}]},
    {'name':'用药','items':[
        {'id':'onmed','type':'single','text':'是否在用药','options':YN,'required':True},
        {'id':'meds','type':'table_mixed','text':'用药清单','dynamic':True,'max_rows':5,
         'columns':[{'id':'d','label':'药名','type':'text','required':True},
                    {'id':'n','label':'日剂量','type':'number','min':0,'max':5000},
                    {'id':'f','label':'频次','type':'select','options':[{'label':'qd','value':1},{'label':'bid','value':2}]}]},
        {'id':'labs','type':'table_input','text':'检验值','rows':['血常规','肝功'],
         'columns':[{'id':'v','label':'结果','type':'text'},{'id':'dd','label':'检验日期','type':'date'}]}]}],
    'logic':[
        {'when':{'field':'onmed','op':'eq','value':0},'then':{'action':'hide','targets':['meds']}},
        {'when':{'field':'age','op':'gte','value':65},'then':{'action':'require','targets':['desc']}},
        {'when':{'all':[{'field':'sex','op':'eq','value':1},{'field':'sym','op':'contains','value':2}]},
         'then':{'action':'show','targets':['dept']}},
        {'action':'exclusive','targets':['desc','idc'],'severity':'warn','message':'演示互斥'}]}
check('定义结构合法', hs.validate_crf_definition(DEF)==[], hs.validate_crf_definition(DEF))
check('方案点名的 8 种基本题型全支持',
      set(hs.CRF_BASIC_TYPES)=={'note','text','paragraph','number','date','single','multi','select'})
check('方案点名的 4 种表格题型全支持', len(hs.CRF_TABLE_TYPES)==4, list(hs.CRF_TABLE_TYPES))
check('铺平后共 12 道题', len(items(DEF))==12, len(items(DEF)))
sub('定义层面的错误要拦住')
bad=hs.validate_crf_definition
check('重复 id 被拦', any('重复' in e for e in bad({'items':[{'id':'a','type':'text','text':'x'},{'id':'a','type':'text','text':'y'}]})))
check('未知题型被拦', any('type' in e for e in bad({'items':[{'id':'a','type':'bogus','text':'x'}]})))
check('选择题没选项被拦', any('options' in e for e in bad({'items':[{'id':'a','type':'single','text':'x'}]})))
check('选项 value 重复被拦', any('value 重复' in e for e in bad({'items':[{'id':'a','type':'single','text':'x','options':[{'label':'p','value':1},{'label':'q','value':1}]}]})))
check('逻辑指向不存在的题被拦', any('不存在' in e for e in bad({'items':[{'id':'a','type':'text','text':'x'}],'logic':[{'when':{'field':'a','op':'eq','value':1},'then':{'action':'hide','targets':['ZZZ']}}]})))
check('条件引用不存在的题被拦', any('不存在' in e for e in bad({'items':[{'id':'a','type':'text','text':'x'}],'logic':[{'when':{'field':'ZZZ','op':'eq','value':1},'then':{'action':'hide','targets':['a']}}]})))
check('下拉框表格里放数字列被拦(该类型只允许 select)', any('只允许' in e for e in bad({'items':[{'id':'t','type':'table_dropdown','text':'x','rows':['r1'],'columns':[{'id':'c','label':'c','type':'number'}]}]})))
check('表格既无固定行也没标 dynamic 被拦', any('永远是空的' in e for e in bad({'items':[{'id':'t','type':'table_input','text':'x','columns':[{'id':'c','label':'c','type':'text'}]}]})))

section('2. 逻辑引擎 (方案 §2.1(2))')
check('在用药时用药清单可见', hs.eval_crf_logic(DEF,{'onmed':1,'age':40})['visible']['meds'] is True)
check('不在用药时用药清单被隐藏', hs.eval_crf_logic(DEF,{'onmed':0,'age':40})['visible']['meds'] is False)
sub('隐藏的必填项不能挡住提交 (最容易踩的坑)')
D2={'items':[{'id':'a','type':'single','text':'有无','options':YN,'required':True},
             {'id':'b','type':'text','text':'详情','required':True}],
    'logic':[{'when':{'field':'a','op':'eq','value':0},'then':{'action':'hide','targets':['b']}}]}
e,w=hs.validate_crf_data(D2,{'a':0})
check('b 被隐藏后不再报"必填未答"', e==[], e)
e,w=hs.validate_crf_data(D2,{'a':1})
check('b 可见时仍然报必填', any(x['field']=='b' for x in e))
sub('隐藏题的旧答案不能驱动后续逻辑(幽灵数据)')
D3={'items':[{'id':'q1','type':'single','text':'A','options':YN},
             {'id':'q2','type':'single','text':'B','options':YN},
             {'id':'q3','type':'text','text':'C','hidden':True}],
    'logic':[{'when':{'field':'q1','op':'eq','value':0},'then':{'action':'hide','targets':['q2']}},
             {'when':{'field':'q2','op':'eq','value':1},'then':{'action':'show','targets':['q3']}}]}
st=hs.eval_crf_logic(D3,{'q1':1,'q2':1})
check('q1=是 时 q2 可见, q2=是 让 q3 显示', st['visible']['q2'] and st['visible']['q3'])
st=hs.eval_crf_logic(D3,{'q1':0,'q2':1,'q3':'残留内容'})
check('q1 改成否 -> q2 隐藏', st['visible']['q2'] is False)
check('q2 的旧答案不再驱动 q3, q3 也随之隐藏', st['visible']['q3'] is False, st['visible'])
check('残留答案被标为 stale 而不是静默丢弃', 'q3' in st['stale'] and 'q2' in st['stale'], st['stale'])
check('lint 对这份写法没有意见(q3 已声明 hidden)', hs.lint_crf_definition(D3)==[])
e,w=hs.validate_crf_data(D3,{'q1':0,'q2':1,'q3':'残留'})
check('stale 只是警告不阻断', e==[] and any(x.get('rule')=='stale_hidden' for x in w))
sub('lint: show 作用在默认可见的题上, 等于没写')
TRAP={'items':[{'id':'a','type':'single','text':'是否在用药','options':YN},
               {'id':'b','type':'text','text':'用药清单'}],
      'logic':[{'when':{'field':'a','op':'eq','value':1},'then':{'action':'show','targets':['b']}}]}
adv=hs.lint_crf_definition(TRAP)
check('结构本身是合法的(不该拦住保存)', hs.validate_crf_definition(TRAP)==[])
check('但 lint 会指出这条 show 等于没写', any(x['kind']=='show_without_default_hidden' for x in adv), adv)
check('并给出改法(加 hidden:true)', adv and 'hidden' in adv[0]['detail'])
check('实测确实一直显示 —— lint 说的是对的', hs.eval_crf_logic(TRAP,{'a':0})['visible']['b'] is True)
NEVER={'items':[{'id':'a','type':'text','text':'永不出现的必填项','hidden':True,'required':True}],'logic':[]}
check('必填+默认隐藏+无人揭开 也会被 lint 指出',
      any(x['kind']=='required_but_never_shown' for x in hs.lint_crf_definition(NEVER)))
sub('其余动作')
check('年龄≥65 时主诉被置为必填', hs.eval_crf_logic(DEF,{'onmed':1,'age':70})['required']['desc'] is True)
check('年龄 30 时主诉非必填', hs.eval_crf_logic(DEF,{'onmed':1,'age':30})['required']['desc'] is False)
check('all 条件(男 且 症状含失眠)成立', hs.eval_crf_logic(DEF,{'onmed':1,'sex':1,'sym':[1,2]})['visible']['dept'] is True)
st=hs.eval_crf_logic(DEF,{'onmed':1,'desc':'头痛','idc':'110101199003077213'})
check('互斥同时填被抓到', any(v['rule']=='exclusive' for v in st['violations']))
D4={'items':[{'id':'a','type':'single','text':'A','options':YN},{'id':'b','type':'number','text':'B'}],
    'logic':[{'when':{'field':'a','op':'eq','value':1},'then':{'action':'set_value','targets':['b'],'value':99}}]}
check('自动设值生效', hs.eval_crf_logic(D4,{'a':1})['auto'].get('b')==99)
D5={'items':[{'id':'a','type':'number','text':'A'},{'id':'b','type':'number','text':'B'}],
    'logic':[{'when':{'all':[{'field':'a','op':'filled'},{'field':'b','op':'filled'}]},
              'then':{'action':'check','targets':['b']},'severity':'block','message':'演示: a 和 b 不该同时有值'}]}
e,w=hs.validate_crf_data(D5,{'a':1,'b':2})
check('逻辑校验 severity=block 进 errors', any('不该同时' in x['error'] for x in e))
D5['logic'][0]['severity']='warn'
e,w=hs.validate_crf_data(D5,{'a':1,'b':2})
check('severity=warn 进 warnings 不阻断', e==[] and w)
sub('规则打架必须收敛, 不能死循环')
OSC={'items':[{'id':'a','type':'text','text':'A'},{'id':'b','type':'text','text':'B'}],
     'logic':[{'when':{'field':'a','op':'empty'},'then':{'action':'hide','targets':['b']}},
              {'when':{'field':'b','op':'empty'},'then':{'action':'hide','targets':['a']}}]}
check('互相隐藏的规则不死循环', hs.eval_crf_logic(OSC,{})['passes']<=4)

section('3. 数据层校验')
GOOD={'nm':'张三','dt':'2026-08-01','onmed':1,'age':45,'sex':1,'sym':[1,3],'dept':1,
      'idc':'110101199003077213','meds':[{'d':'氯氮平','n':100,'f':1}],
      'labs':[{'v':'正常','dd':'2026-07-30'},{'v':'轻度异常','dd':'2026-07-30'}]}
check('一份合规填报无错误', hs.validate_crf_data(DEF,GOOD)[0]==[], hs.validate_crf_data(DEF,GOOD)[0])
for mut,label in [({'age':200},'数字超范围'),({'dt':'2026/08/01'},'日期格式错'),
                  ({'sex':9},'单选值不在选项内'),({'sym':1},'多选答案不是数组'),
                  ({'sym':[1,99]},'多选含非法值'),({'idc':'110101199003077219'},'身份证校验位')]:
    e,_=hs.validate_crf_data(DEF, dict(GOOD, **mut))
    check(label+'被抓', bool(e), [x['error'] for x in e][:2])
sub('表格题')
for mut,kw in [({'meds':[{'d':'','n':100,'f':1}]},'必填'),
               ({'meds':[{'d':'x','n':99999,'f':1}]},'超出范围'),
               ({'meds':[{'d':'x','zz':1}]},'未定义的列'),
               ({'labs':[{'v':'a','dd':'2026-07-30'}]},'固定行'),
               ({'meds':[{'d':str(i)} for i in range(9)]},'最多'),
               ({'labs':[{'v':'a','dd':'30/07/2026'},{'v':'b'}]},'YYYY-MM-DD')]:
    e,_=hs.validate_crf_data(DEF, dict(GOOD, **mut))
    check(kw+' 被抓', any(kw in x['error'] for x in e), [x['error'] for x in e][:2])

section('4. 版本管理: "修改不影响已有数据" (方案 §2.1(3))')
r,err=hs.upsert_platform_crf({'code':PFX+'A','name':'版本测试表','definition':DEF,'category':'演示','owner':'医生甲'})
check('CRF 入库', r and r['version']=='1', err or r)
sub('删掉被逻辑引用的题, 定义校验就该拦住')
D_bad=json.loads(json.dumps(DEF))
D_bad['sections'][0]['items']=[i for i in D_bad['sections'][0]['items'] if i['id']!='desc']
r,err=hs.upsert_platform_crf({'code':PFX+'A','name':'x','definition':D_bad})
check('删掉仍被互斥规则引用的题 -> 拒绝并指名道姓', r is None and err and 'desc' in err, err)
sub('没有填报时: 怎么改都就地改')
D_del=json.loads(json.dumps(DEF))
D_del['sections'][0]['items']=[i for i in D_del['sections'][0]['items'] if i['id']!='nm']
r,err=hs.upsert_platform_crf({'code':PFX+'A','name':'版本测试表','definition':D_del})
check('无填报时删题也不涨版本', r and r['version']=='1', err or (r or {}).get('version'))
check('但仍然告诉你这是破坏性改动', r and '破坏性' in (r.get('note') or ''), (r or {}).get('note'))
hs.upsert_platform_crf({'code':PFX+'A','name':'版本测试表','definition':DEF})
sub('有填报之后')
sub_r,err=hs.submit_crf_response({'crf_code':PFX+'A','patient_no':PFX+'001','data':GOOD,
                                  'operator':'医生甲','allow_warnings':True})
check('填报成功', sub_r and sub_r.get('accepted'), err or sub_r)
check('填报钉住了版本 1', sub_r['crf_version']=='1')
first_id=sub_r['id']
D_safe=json.loads(json.dumps(DEF))
D_safe['sections'][0]['items'][1]['text']='患者姓名(必填)'
D_safe['sections'][0]['items'].append({'id':'newf','type':'text','text':'新增字段'})
D_safe['sections'][0]['items'][3]['max']=150
ch=hs.classify_crf_change(DEF,D_safe)
check('改措辞/新增题/放宽范围 全判为安全', ch['verdict']=='in_place', ch['breaking'])
check('安全改动逐条列了出来', len(ch['safe'])>=3, [x['kind'] for x in ch['safe']])
r,err=hs.upsert_platform_crf({'code':PFX+'A','name':'版本测试表','definition':D_safe})
check('安全改动就地改, 版本仍是 1', r and r['version']=='1')
D_ord=json.loads(json.dumps(DEF)); D_ord['sections'][0]['items'].reverse()
ch=hs.classify_crf_change(DEF,D_ord)
check('调整顺序是安全改动(数据按 id 存)', ch['verdict']=='in_place')
check('顺序变化有记录', any(x['kind']=='reordered' for x in ch['safe']))
for name,mut,kind in [('删题',lambda d:d['sections'][0]['items'].pop(1),'item_removed'),
                      ('改题型',lambda d:d['sections'][0]['items'][3].__setitem__('type','text'),'type_changed'),
                      ('删选项',lambda d:d['sections'][0]['items'][6]['options'].pop(),'option_removed'),
                      ('收紧 max',lambda d:d['sections'][0]['items'][3].__setitem__('max',60),'max_tightened'),
                      ('删表格列',lambda d:d['sections'][1]['items'][2]['columns'].pop(),'column_removed')]:
    d=json.loads(json.dumps(DEF)); mut(d)
    ch=hs.classify_crf_change(DEF,d)
    check(name+' 判为破坏性', ch['verdict']=='new_version' and any(x['kind']==kind for x in ch['breaking']),
          [x['kind'] for x in ch['breaking']])
D_break=json.loads(json.dumps(DEF))
D_break['sections'][1]['items']=[i for i in D_break['sections'][1]['items'] if i['id']!='labs']
r,err=hs.upsert_platform_crf({'code':PFX+'A','name':'版本测试表','definition':D_break})
check('破坏性改动 + 已有填报 => 自动开新版 2', r and r['version']=='2', (r or {}).get('version'))
check('说清了为什么开新版', r and '自动开新版' in (r.get('note') or ''))
q,_=hs.query_platform_crfs(code=PFX+'A', all_versions=True)
check('两个版本并存', q['count']==2, [c['version'] for c in q['crfs']])
v1=[c for c in q['crfs'] if c['version']=='1'][0]
check('版本 1 的定义没被动过(被删的检验值表还在)', any(i['id']=='labs' for i in items(v1['definition'])))
v2=[c for c in q['crfs'] if c['version']=='2'][0]
check('版本 2 里确实删掉了', not any(i['id']=='labs' for i in items(v2['definition'])))
check('版本 1 记着它有 1 份填报', v1['response_count']==1)
check('列表页题数不依赖 definition 也算得对(分节表单)',
      all(c['item_count'] is not None for c in hs.query_platform_crfs(all_versions=True)[0]['crfs']))
rs,_=hs.query_crf_responses(patient_no=PFX+'001')
check('旧填报仍钉在版本 1 上', rs['responses'][0]['crf_version']=='1')
check('旧填报里被删字段的数据仍完整可读', rs['responses'][0]['data'].get('labs')==GOOD['labs'])
check('版本号自增: 1->2', hs._bump_version('1')=='2')
check('版本号自增: v1.2->v1.3', hs._bump_version('v1.2')=='v1.3')
check('认不出数字时挂 -2', hs._bump_version('rev')=='rev-2')

section('5. 提交行为')
s1,err=hs.submit_crf_response({'crf_code':PFX+'A','crf_version':'1','patient_no':PFX+'001','data':dict(GOOD,nm='')})
check('必填缺失被拒收', s1 and s1.get('accepted') is False and s1['errors'])
s2,err=hs.submit_crf_response({'crf_code':PFX+'A','crf_version':'1','patient_no':PFX+'001','data':dict(GOOD,onmed=0)})
check('有弱校验时默认不收, 并给出提示怎么办', s2 and s2.get('accepted') is False and s2.get('hint'), (s2 or {}).get('hint'))
s3,err=hs.submit_crf_response({'crf_code':PFX+'A','crf_version':'1','patient_no':PFX+'001',
                               'data':dict(GOOD,onmed=0),'allow_warnings':True})
check('确认后可提交', s3 and s3.get('accepted'), err)
check('被隐藏题的残留答案单独归档, 不进正式数据', s3['archived_hidden']==1, s3)
rs,_=hs.query_crf_responses(patient_no=PFX+'001', limit=1)
rec=rs['responses'][0]
check('正式数据里没有用药清单', 'meds' not in rec['data'])
check('残留答案在 hidden_data 里可查', 'meds' in (rec['hidden_data'] or {}))
sub('自动设值以规则为准, 不听客户端的')
hs.upsert_platform_crf({'code':PFX+'B','name':'自动设值表','definition':D4})
s4,err=hs.submit_crf_response({'crf_code':PFX+'B','patient_no':PFX+'001','data':{'a':1,'b':7}})
check('提交成功', s4 and s4.get('accepted'), err)
rs,_=hs.query_crf_responses(code=PFX+'B')
check('客户端传的 b=7 被规则算出的 99 覆盖', rs['responses'][0]['data'].get('b')==99, rs['responses'][0]['data'])
sub('修订留痕')
s5,_=hs.submit_crf_response({'crf_code':PFX+'A','crf_version':'1','patient_no':PFX+'001',
                             'data':dict(GOOD,nm='张三丰'),'revision_of':first_id,'allow_warnings':True})
check('修订成功', s5 and s5.get('accepted'))
rs,_=hs.query_crf_responses(patient_no=PFX+'001', include_superseded=True)
old=[x for x in rs['responses'] if x['id']==first_id][0]
check('被修订的旧版标成 superseded', old['status']=='superseded')
check('默认列表不含已被取代的版本',
      all(x['id']!=first_id for x in hs.query_crf_responses(patient_no=PFX+'001')[0]['responses']))

section('6. 拷贝 / 私有与共享 (方案 §2.1(3))')
r,err=hs.copy_platform_crf({'code':PFX+'A','new_code':PFX+'C','new_name':'拷贝出来的表',
                            'scope':'shared','owner':'医生乙'})
check('拷贝成功', r and r['code']==PFX+'C', err)
check('拷贝出来是新表的第 1 版, 不是原表的新版本', r['version']=='1')
q,_=hs.query_platform_crfs(code=PFX+'C')
check('记着拷贝自哪一版', (q['crfs'][0]['copied_from'] or '').startswith(PFX+'A@'))
check('scope 可设为共享', q['crfs'][0]['scope']=='shared')
check('不许拷成同名', hs.copy_platform_crf({'code':PFX+'A','new_code':PFX+'A'})[0] is None)
check('目标 code 已存在时拒绝', hs.copy_platform_crf({'code':PFX+'A','new_code':PFX+'C'})[0] is None)
check('源不存在时拒绝', hs.copy_platform_crf({'code':'NOPE','new_code':PFX+'D'})[0] is None)

section('7. AI 生成 CRF (方案 §2.1(1))')
d,rp=hs.generate_crf_draft({'disease':'2型糖尿病','visit_type':'随诊',
                            'fields':['末次糖化血红蛋白','是否发生低血糖','足部检查日期','并发症情况']})
check('生成通过结构校验', rp['validation']==[], rp['validation'])
check('生成了多个章节', rp['section_count']>=4, rp['section_count'])
check('生成了逻辑规则', rp['logic_count']>=3)
types={i['type'] for i in items(d['definition'])}
check('覆盖了多种题型', len(types)>=6, sorted(types))
check('含表格题', any(t in hs.CRF_TABLE_TYPES for t in types))
fld={i['text']:i['type'] for i in items(d['definition'])}
check('"…日期"猜成日期题', fld.get('足部检查日期')=='date')
check('"是否…"猜成单选题', fld.get('是否发生低血糖')=='single')
check('"…情况"猜成段落题', fld.get('并发症情况')=='paragraph')
check('猜不准的说明里明写了是猜的', any('猜' in n['detail'] for n in rp['notes'] if n['step']=='custom_fields'))
d2,_=hs.generate_crf_draft({'disease':'2型糖尿病','visit_type':'随诊'})
check('默认 code 稳定可复现(不用内置 hash)', d['code']==d2['code'], (d['code'],d2['code']))
st=hs.eval_crf_logic(d['definition'],{'on_med':0,'has_ae':0})
check('生成的逻辑真的能跑: 不用药则隐藏用药清单', st['visible']['med_table'] is False)
check('生成的逻辑真的能跑: 无不良事件则隐藏描述', st['visible']['ae_desc'] is False)
check('有不良事件时描述被置为必填', hs.eval_crf_logic(d['definition'],{'on_med':1,'has_ae':1})['required']['ae_desc'] is True)

section('8. Excel 建表 (方案 §2.1(4))')
try:
    import openpyxl
    HAVE=True
except ImportError:
    HAVE=False
    check('openpyxl 缺失时给出可执行的提示而不是崩溃',
          'pip install openpyxl' in (hs.parse_excel_to_crf(b'x')[2] or ''))
if HAVE:
    sub('A: 配置式表格')
    wb=openpyxl.Workbook(); ws=wb.active; ws.title='随访CRF'
    ws.append(['变量名','题干','题型','选项','必填','章节'])
    for row in [['nm','姓名','文本','','是','基本信息'],['age','年龄','数值','','','基本信息'],
                ['sex','性别','单选','男;女','是','基本信息'],['dt','入组日期','日期','','','基本信息'],
                ['bad','这题是单选但没给选项','单选','','','其他']]:
        ws.append(row)
    buf=io.BytesIO(); wb.save(buf)
    d,rp,err=hs.parse_excel_to_crf(buf.getvalue())
    check('配置式解析成功', d and not err, err)
    check('识别为配置式', rp['mode']=='config')
    check('解析出 5 道题', rp['item_count']==5)
    check('结构校验通过', rp['validation']==[], rp['validation'])
    t={i['text']:i for i in items(d['definition'])}
    check('文本/数值/单选/日期都认出来了',
          [t['姓名']['type'],t['年龄']['type'],t['性别']['type'],t['入组日期']['type']]==['text','number','single','date'])
    check('选项被切开', len(t['性别']['options'])==2)
    check('必填列被识别', t['姓名'].get('required') is True)
    check('没给选项的单选降级为文本并说明',
          t['这题是单选但没给选项']['type']=='text' and any('降级' in n['detail'] for n in rp['notes']))
    check('分章节了', len(d['definition'].get('sections') or [])==2)
    sub('B: 数据式表格(靠实际取值推题型)')
    wb=openpyxl.Workbook(); ws=wb.active; ws.title='历史随访表'
    ws.append(['门诊号','就诊日期','收缩压','转归','医生小结'])
    for i,(no,dt,bp,oc) in enumerate([('P001','2026-01-05',128,'好转'),('P002','2026-01-06',145,'稳定'),
                                      ('P003','2026-01-07',132,'好转'),('P004','2026-01-08',150,'恶化'),
                                      ('P005','2026-01-09',121,'稳定'),('P006','2026-01-10',139,'好转')]):
        ws.append([no,dt,bp,oc,'患者本次随访情况总体平稳，建议继续原方案治疗并两周后复诊复查血压'+str(i)])
    buf=io.BytesIO(); wb.save(buf)
    d,rp,err=hs.parse_excel_to_crf(buf.getvalue())
    check('数据式解析成功', d and not err, err)
    check('识别为数据式', rp['mode']=='data')
    t={i['text']:i for i in items(d['definition'])}
    check('整列日期 -> 日期题', t['就诊日期']['type']=='date')
    check('整列数字 -> 数字题', t['收缩压']['type']=='number')
    check('少量重复取值 -> 单选题', t['转归']['type']=='single')
    check('单选选项来自实际取值', len(t['转归'].get('options') or [])==3)
    check('长文本 -> 段落题', t['医生小结']['type']=='paragraph')
    check('段落判定说明里给了中位数依据',
          any('中位数' in n['detail'] and '医生小结' in n['detail'] for n in rp['notes']))
    check('零散短文本 -> 文本题', t['门诊号']['type']=='text')
    check('数字题给了实测 min/max', t['收缩压'].get('min')==121 and t['收缩压'].get('max')==150)
    check('并说明了那只是样本范围不是业务约束', any('样本范围' in n['detail'] for n in rp['notes']))
    check('每一列都写明了凭什么这么判', sum(1 for n in rp['notes'] if n['step']=='column_type')==5)
    sub('一列短代码里混一条长备注, 不该整列判成段落')
    wb=openpyxl.Workbook(); ws=wb.active; ws.title='混合列'
    ws.append(['科室'])
    for v in ['内科','外科','儿科','内科','外科','本例因患者依从性差且合并多种基础疾病故转多学科联合门诊长期随访管理']:
        ws.append([v])
    buf=io.BytesIO(); wb.save(buf)
    d2,rp2,_=hs.parse_excel_to_crf(buf.getvalue())
    t2={i['text']:i for i in items(d2['definition'])}
    check('混了一条长值仍判为短取值类型(不是段落)', t2['科室']['type']!='paragraph', t2['科室']['type'])
    check('坏文件给明确报错而不是抛异常', hs.parse_excel_to_crf(b'not an xlsx')[2])

section('9. 清理')
clean()
print('  ✅ 测试数据已清')
finish()
